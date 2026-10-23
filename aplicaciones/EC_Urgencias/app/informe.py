# -*- coding: utf-8 -*-
"""Generación del informe Word (.docx) con las tablas y gráficas de todos los escenarios."""
import datetime
import io
import math

import numpy as np
import pandas as pd
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from motor import TIPOS
from kpis import (ETIQUETAS_KPI, KPIS_SINTESIS, num, pct, eur, fmt_kpi, sentido_kpi, resumen_mc, tabla_media_mc,
                  slas_mc, comparar_escenarios, viabilidad_economica, tabla_sintesis, tabla_economica,
                  tabla_slas, texto_escenario, texto_recomendacion)
import graficos as gr

def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


class Informe:
    """Envoltorio sencillo de python-docx con helpers para títulos, párrafos, tablas y figuras."""

    def __init__(self, titulo, subtitulo=""):
        self.doc = docx.Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Calibri"; estilo.font.size = Pt(10.5)
        self.n_tabla, self.n_fig = 0, 0
        h = self.doc.add_heading(titulo, level=0)
        if subtitulo:
            p = self.doc.add_paragraph(subtitulo); p.runs[0].italic = True

    def h(self, texto, nivel=1):
        self.doc.add_heading(texto, level=nivel)

    def p(self, texto, negrita_inicio=None):
        par = self.doc.add_paragraph()
        if negrita_inicio == "auto":      # negrita hasta los primeros ':' o '. '
            i, j = texto.find(": "), texto.find(". ")
            cortes = [x for x in (i, j) if x > 0]
            negrita_inicio = texto[:min(cortes) + 1] if cortes and min(cortes) < 60 else None
        if negrita_inicio and texto.startswith(negrita_inicio):
            par.add_run(negrita_inicio).bold = True
            par.add_run(texto[len(negrita_inicio):])
        else:
            par.add_run(texto)
        par.paragraph_format.space_after = Pt(6)
        return par

    def vinetas(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def tabla(self, df, titulo, formato=None, indice=True, ancho_fuente=8, nombre_indice=""):
        """Inserta un DataFrame como tabla. `formato`: función (fila, columna, valor) -> str."""
        self.n_tabla += 1
        cap = self.doc.add_paragraph()
        cap.add_run(f"Tabla {self.n_tabla}. {titulo}").bold = True
        d = df.copy()
        cols = list(d.columns)
        ncols = len(cols) + (1 if indice else 0)
        t = self.doc.add_table(rows=1, cols=ncols)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cab = ([nombre_indice or (d.index.name or "")] if indice else []) + [str(c) for c in cols]
        for j, c in enumerate(cab):
            cel = t.rows[0].cells[j]; cel.text = ""
            r = cel.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(ancho_fuente)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _sombrear(cel, "2F5496")
        for i, (idx, fila) in enumerate(d.iterrows()):
            celdas = t.add_row().cells
            vals = ([str(idx)] if indice else []) + [
                formato(idx, c, fila[c]) if formato else _fmt_generico(fila[c]) for c in cols]
            crudos = ([idx] if indice else []) + [fila[c] for c in cols]
            for j, v in enumerate(vals):
                celdas[j].text = ""
                r = celdas[j].paragraphs[0].add_run(v); r.font.size = Pt(ancho_fuente)
                if (j > 0 or not indice) and not isinstance(crudos[j], str):
                    celdas[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i % 2 == 1:
                    _sombrear(celdas[j], "EAF0F8")
        self.doc.add_paragraph()

    def figura(self, buf, titulo, ancho_cm=16.5):
        if buf is None:
            return
        if not isinstance(buf, io.BytesIO):          # se admite directamente una Figure
            buf = fig_a_png(buf)
        self.n_fig += 1
        self.doc.add_picture(buf, width=Cm(ancho_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figura {self.n_fig}. {titulo}"); r.italic = True; r.font.size = Pt(9)

    def salto(self):
        self.doc.add_page_break()

    def guardar(self, ruta=None):
        """Guarda en `ruta` o, si no se indica, devuelve el documento como bytes."""
        if ruta:
            self.doc.save(ruta)
            return ruta
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def _fmt_generico(v):
    if isinstance(v, (bool, np.bool_)):
        return "Sí" if v else "No"
    if isinstance(v, (int, np.integer)):
        return num(v, 0)
    if isinstance(v, (float, np.floating)):
        if pd.isna(v):
            return "—"
        return num(v, 0) if abs(v) >= 1000 else num(v, 3)
    return str(v)


fig_a_png = gr.fig_a_png



# ----------------------------------------------------------------------------------------------
# Utilidades de formato
# ----------------------------------------------------------------------------------------------
def fmt_por_fila(fila, col, v):
    """Formato según el KPI de la fila (tablas con KPIs en filas)."""
    if isinstance(v, str):
        return v
    clave = {vv: kk for kk, vv in ETIQUETAS_KPI.items()}.get(fila, fila)
    if col == "p_valor":
        return num(v, 4)
    if col == "Δ%":
        return pct(v, 1)
    if col == "Significativo":
        return "Sí" if v else "No"
    return fmt_kpi(str(clave), v)


def _aplanar(d, prefijo=""):
    """Aplana diccionarios anidados: {'a': {'b': 1}} -> {'a › b': 1}."""
    out = {}
    for k, v in d.items():
        clave = f"{prefijo} › {k}" if prefijo else str(k)
        if isinstance(v, dict):
            out.update(_aplanar(v, clave))
        else:
            out[clave] = v
    return out


ATRIBUTOS_ESCENARIO = ["SERVIDORES_CAPACIDAD", "CAMAS_CAPACIDAD", "CAMA_ALIAS", "POOL_ESTACION", "BOX_CAPACIDAD",
                       "BOX_CAPTURA", "UCI_MODE", "UCI_ROJO_CAPACIDAD", "UCI_NARANJA_CAPACIDAD", "UCI_BASE_CAPACIDAD",
                       "UCI_TOTAL_CAPACIDAD", "LIMITE_ESPERA_UCI", "TIEMPOS_SERVICIO", "PROB_DERIVACION", "COSTES_FIJOS"]


def diferencias_config(cfg_ref, cfg):
    """Tabla de parámetros que difieren entre dos configuraciones."""
    a = _aplanar({k: getattr(cfg_ref, k, None) for k in ATRIBUTOS_ESCENARIO})
    b = _aplanar({k: getattr(cfg, k, None) for k in ATRIBUTOS_ESCENARIO})
    filas = []
    for k in sorted(set(a) | set(b)):
        va, vb = a.get(k), b.get(k)
        if isinstance(va, (int, float)) and isinstance(vb, (int, float)) and math.isclose(va, vb, rel_tol=1e-9):
            continue
        if va != vb:
            filas.append({"Parámetro": k, "Referencia": "—" if va is None else _txt(va, k),
                          "Escenario": "—" if vb is None else _txt(vb, k)})
    return pd.DataFrame(filas)


def _txt(v, clave=""):
    if clave.startswith("COSTES_FIJOS") and isinstance(v, (int, float)):
        return eur(v)
    if isinstance(v, float):
        if clave.startswith("TIEMPOS_SERVICIO"):
            return f"{num(v * 60, 1)} min" if v < 2 else f"{num(v, 1)} h"
        return num(v, 3)
    return str(v)


def tabla_capacidades(configs):
    filas = {}
    for et, c in configs.items():
        col = {**{f"Servidor · {k}": v for k, v in c.SERVIDORES_CAPACIDAD.items()},
               **{f"Cama · {k}": v for k, v in c.CAMAS_CAPACIDAD.items()},
               "Box de Urgencias": c.BOX_CAPACIDAD,
               "UCI · modo": "Bolsa compartida (prioridad)" if c.UCI_MODE == "pooled" else "Bolsas reservadas",
               "UCI · camas totales": c.UCI_TOTAL_CAPACIDAD,
               "UCI · Rojo / Naranja / Base": ("—" if c.UCI_MODE == "pooled" else
                                               f"{c.UCI_ROJO_CAPACIDAD} / {c.UCI_NARANJA_CAPACIDAD} / {c.UCI_BASE_CAPACIDAD}"),
               "Límite de espera UCI (h)": c.LIMITE_ESPERA_UCI,
               "Coste fijo anual (€)": sum(v for k, v in c.COSTES_FIJOS.items() if "anual" in k),
               "Coste fijo único (€)": sum(v for k, v in c.COSTES_FIJOS.items() if "anual" not in k)}
        filas[et] = col
    return pd.DataFrame(filas)


KPIS_GRAFICOS = ["T_sys_media_Verde", "T_sys_media_Amarillo", "T_sys_media_Global", "Latencia_DI_media_NR_min",
                 "Tasa_mortalidad_Global", "Mort_Sin_UCI_NR", "Ocup_media_UCI", "Beneficio_operativo_dia", "Margen_neto"]


# ----------------------------------------------------------------------------------------------
# Informe
# ----------------------------------------------------------------------------------------------
def generar_informe(protocolos, ref, ajustes, comunes, titulo="Red de colas hospitalaria · Servicio de Urgencias",
                    progreso=None):
    """Construye el informe y lo devuelve como bytes.
    protocolos: {nombre: resultado de montecarlo}; ref: nombre del escenario de referencia;
    ajustes: dict con n_sim, dias, semilla, warmup, cooldown; comunes: dict con 'llegadas' (lista) y 'param'."""
    pasos = 6 + len(protocolos)
    paso = [0]

    def avanza(txt):
        paso[0] += 1
        if progreso:
            progreso(paso[0] / pasos, txt)

    param = comunes["param"]
    inf = Informe(titulo, f"Informe de simulación generado el {datetime.date.today().strftime('%d/%m/%Y')} · "
                          f"{ajustes['n_sim']} réplicas de {ajustes['dias']} días por escenario · semilla inicial "
                          f"{ajustes['semilla']} · referencia: {ref}")
    configs = {et: m["cfg"] for et, m in protocolos.items()}
    orden = list(protocolos)

    # --- Resumen ejecutivo
    avanza("Resumen ejecutivo")
    recomendacion, ranking = texto_recomendacion(protocolos, ref)
    inf.h("Resumen ejecutivo", 1)
    for p in recomendacion:
        inf.p(p, negrita_inicio="auto")
    inf.figura(gr.fig_sintesis(ranking, orden, ref), "Resultado neto, SLAs cumplidos y mortalidad Sin_UCI por escenario")

    # --- Configuración de la simulación y parámetros comunes
    avanza("Parámetros")
    inf.salto(); inf.h("1. Configuración de la simulación", 1)
    inf.tabla(pd.DataFrame({"Valor": {
        "Nº de réplicas por escenario": ajustes["n_sim"], "Días por réplica": ajustes["dias"],
        "Semilla inicial (réplica i → semilla + i)": ajustes["semilla"],
        "Periodo de calentamiento descartado (días)": ajustes["warmup"],
        "Periodo final descartado (días)": ajustes["cooldown"],
        "Escenario de referencia": ref, "Escenarios simulados": ", ".join(orden)}}),
        "Parámetros de la simulación", nombre_indice="Parámetro", formato=lambda f, c, v: str(v))
    inf.h("2. Parámetros comunes a todos los escenarios", 1)
    lleg = pd.DataFrame(comunes["llegadas"])
    lleg.index = [f"{int(a):02d}:00–{int(b):02d}:00" for a, b in zip(lleg["Inicio"], lleg["Fin"])]
    inf.tabla(lleg.drop(columns=["Inicio", "Fin"]), "Llegadas: tasa (pacientes/hora) y reparto por perfil", nombre_indice="Franja",
              formato=lambda f, c, v: num(v, 2) if c == "Tasa" else pct(v, 0))
    eco = pd.DataFrame({"Reembolso por alta (€)": pd.Series(param["REEMBOLSO"])}).reindex(TIPOS)
    inf.tabla(eco, "Reembolso por alta según perfil", nombre_indice="Perfil", formato=lambda f, c, v: eur(v))
    costes = pd.concat([pd.Series({f"{k} (€/servicio)": v for k, v in param["COSTE_SERVICIO"].items()}),
                        pd.Series({f"{k} (€/día)": v for k, v in param["COSTE_DIA"].items()}),
                        pd.Series({"Penalización por fallecimiento (€)": param["PEN_FALLECE"],
                                   "Penalización adicional Sin_UCI (€)": param["PEN_SIN_UCI"],
                                   "Penalización por demora crítica N/R (€)": param.get("PEN_DEMORA", 0)})]).to_frame("Importe")
    inf.tabla(costes, "Costes variables y penalizaciones", nombre_indice="Concepto", formato=lambda f, c, v: eur(v))
    tramos = pd.DataFrame([{"Tasa de altas mayor que": pct(u, 0), "Margen exigido": pct(m, 1)} for u, m in param["TRAMOS_MARGEN"]]
                          + [{"Tasa de altas mayor que": "resto", "Margen exigido": pct(param["MARGEN_MINIMO"], 1)}])
    inf.tabla(tramos, "Margen operativo exigido por el convenio", indice=False, formato=lambda f, c, v: str(v))
    inf.tabla(pd.DataFrame({"Umbral": {
        "SLA1 · Latencia máx. hasta Diag. Inicial (Naranja/Rojo)": f"{num(param['SLA_LATENCIA_H'] * 60, 0)} min en ≥{pct(param['SLA_LATENCIA_CUMPL'], 0)} de pacientes",
        "SLA2 · Ratio de espera Verde/Amarillo": f"< {pct(param['SLA_RATIO_ESPERA'], 0)}",
        "SLA3 · Ocupación media de UCI": f"≤ {pct(param['SLA_OCUP_UCI'], 0)}",
        "SLA4 · Mortalidad Sin_UCI (Naranja/Rojo)": f"< {pct(param['SLA_SIN_UCI'], 1)}",
        "SLA5 · Margen neto": "≥ margen exigido"}}), "Acuerdos de nivel de servicio", nombre_indice="SLA",
        formato=lambda f, c, v: str(v))

    inf.h("3. Parámetros específicos de cada escenario", 1)
    inf.tabla(tabla_capacidades(configs), "Recursos, capacidades y costes fijos por escenario", nombre_indice="Parámetro",
              ancho_fuente=7, formato=lambda f, c, v: (eur(v) if "€" in f else (num(v, 0) if isinstance(v, (int, np.integer)) else (num(v, 1) if isinstance(v, float) else str(v)))))
    for et, c in configs.items():
        if et == ref:
            continue
        dif = diferencias_config(configs[ref], c)
        if getattr(c, "descripcion", ""):
            inf.p(f"{et}: {c.descripcion}", negrita_inicio="auto")
        if len(dif):
            inf.tabla(dif.set_index("Parámetro"), f"Parámetros de {et} que difieren de {ref}", nombre_indice="Parámetro",
                      ancho_fuente=7, formato=lambda f, cc, v: str(v))
        else:
            inf.p(f"{et} no difiere de {ref} en ningún parámetro.")

    # --- Comparativa
    avanza("Comparativa")
    inf.salto(); inf.h("4. Comparativa de escenarios", 1)
    medias, marcas = tabla_sintesis(protocolos, ref=ref)
    tab = medias.copy().astype(object)
    for c in tab.columns:
        tab[c] = [fmt_kpi(k, medias.loc[k, c]) + (f" {marcas.loc[k, c]}" if marcas.loc[k, c] else "") for k in medias.index]
    tab = tab.rename(index=ETIQUETAS_KPI)
    for i in range(0, len(orden), 5):          # máximo 5 escenarios por tabla para que quepa en la página
        cols = orden[i:i + 5]
        inf.tabla(tab[cols], f"KPIs principales por escenario (media de las réplicas; ▲ mejora / ▼ empeora "
                             f"significativamente frente a {ref})", nombre_indice="KPI", ancho_fuente=7.5,
                  formato=lambda f, c, v: v)
    inf.tabla(tabla_slas(protocolos).rename(index=slas_mc(protocolos[ref])["Descripción"]),
              "Cumplimiento de SLAs (% de réplicas que lo cumplen)", nombre_indice="SLA", ancho_fuente=7,
              formato=lambda f, c, v: pct(v, 0))
    if len(orden) > 1:
        eco = tabla_economica(protocolos, ref=ref)
        inf.tabla(eco[["Δ beneficio operativo anual (€)", "IC95 inf (€)", "IC95 sup (€)", "Coste fijo anual (€)",
                       "Coste fijo único (€)", "Resultado neto 1er año (€)", "Resultado neto años sucesivos (€/año)"]],
                  f"Resultado económico frente a {ref}", nombre_indice="Escenario", ancho_fuente=7, formato=lambda f, c, v: eur(v))
    inf.tabla(ranking[["SLAs cumplidos", "Resultado neto anual (€)", "Resultado neto 1er año (€)", "Mortalidad global",
                       "Mortalidad Sin_UCI N/R", "T_sys global (h)", "Margen neto", "Descartado"]],
              "Ranking de escenarios (reglas: no empeorar mortalidad → más SLAs → mayor resultado neto)",
              nombre_indice="Escenario", ancho_fuente=7,
              formato=lambda f, c, v: (eur(v) if "€" in c else (pct(v, 2) if ("Mortalidad" in c or "Margen" in c)
                                                               else _fmt_generico(v))))
    inf.figura(gr.fig_comparativo(protocolos, KPIS_GRAFICOS, 3), "Distribución Monte Carlo de los KPIs clave por escenario", 17)
    inf.figura(gr.fig_utilizacion_comparada(protocolos), "Utilización media de recursos por escenario")
    inf.figura(gr.fig_tsys_perfil(protocolos), "Tiempo en el sistema por perfil y escenario")
    inf.figura(gr.fig_ocupacion_uci(protocolos, param["SLA_OCUP_UCI"]), "Ocupación diaria de la UCI por escenario")

    # --- Detalle por escenario
    for n, (et, mc) in enumerate(protocolos.items(), start=1):
        avanza(f"Escenario {et}")
        inf.salto(); inf.h(f"5.{n}. Escenario: {et}", 1)
        if getattr(mc["cfg"], "descripcion", ""):
            inf.p(mc["cfg"].descripcion)
        if et != ref:
            inf.p(texto_escenario(et, protocolos[ref], mc, ref=ref), negrita_inicio="auto")
        inf.tabla(resumen_mc(mc, KPIS_SINTESIS).rename(index=ETIQUETAS_KPI)[["Media", "SD", "IC95_inf", "IC95_sup", "P2.5", "P97.5"]],
                  f"KPIs principales · {et} (Monte Carlo)", formato=fmt_por_fila, nombre_indice="KPI", ancho_fuente=7)
        inf.tabla(tabla_media_mc(mc, "operativos")[["Pacientes", "T_sys_media", "T_sys_P90", "T_ser_media", "T_q_media",
                                                    "Ratio_espera_medio"]],
                  f"KPIs temporales por perfil (h) · {et}", nombre_indice="Perfil",
                  formato=lambda f, c, v: num(v, 0) if c == "Pacientes" else (pct(v) if "Ratio" in c else num(v, 2)))
        inf.tabla(tabla_media_mc(mc, "clinicos")[["Tasa_alta", "Tasa_mortalidad", "Tasa_mort_Sin_UCI", "Tasa_hospitalizacion",
                                                  "Tasa_UCI", "Tasa_Quirofano"]],
                  f"KPIs clínicos por perfil · {et}", nombre_indice="Perfil", formato=lambda f, c, v: pct(v, 2))
        rec = tabla_media_mc(mc, "recursos")
        inf.tabla(rec[["Tipo", "Capacidad", "Utilizacion", "Lq_media", "%Bloqueo", "Diagnostico"]],
                  f"Recursos: utilización, cola media, bed-blocking y diagnóstico · {et}", nombre_indice="Recurso", ancho_fuente=7,
                  formato=lambda f, c, v: v if isinstance(v, str) else (pct(v) if c in ("Utilizacion", "%Bloqueo")
                                                                        else num(v, 0 if c == "Capacidad" else 3)))
        inf.figura(gr.fig_utilizacion(rec, f"Utilización media de recursos · {et}"), f"Utilización media de recursos · {et}", 14)
        inf.tabla(tabla_media_mc(mc, "economicos")[["Pacientes", "Ingresos", "Coste_total", "Beneficio", "Coste_medio_paciente", "Margen"]],
                  f"Resultado económico por perfil (media por réplica) · {et}", nombre_indice="Perfil", ancho_fuente=7,
                  formato=lambda f, c, v: num(v, 0) if c == "Pacientes" else (pct(v) if c == "Margen" else eur(v)))
        s = slas_mc(mc)
        inf.tabla(s.set_index("Descripción")[["Valor_medio", "IC95_inf", "IC95_sup", "%Réplicas_cumple", "Veredicto"]],
                  f"SLAs · {et}", nombre_indice="SLA", ancho_fuente=7,
                  formato=lambda f, c, v: v if isinstance(v, str) else (pct(v, 0) if c == "%Réplicas_cumple" else num(v, 4)))
        if et != ref:
            cmp_ = comparar_escenarios(protocolos[ref], mc, KPIS_SINTESIS)
            inf.tabla(cmp_[["Base", "Escenario", "Δ%", "p_valor", "Significativo"]].rename(index=ETIQUETAS_KPI)
                      .rename(columns={"Base": ref, "Escenario": et}),
                      f"Comparación pareada {et} vs. {ref}", formato=fmt_por_fila, nombre_indice="KPI", ancho_fuente=7)
            viab = viabilidad_economica(protocolos[ref], mc)
            inf.tabla(viab, f"Viabilidad económica · {et} vs. {ref}", nombre_indice="Concepto", ancho_fuente=7.5,
                      formato=lambda f, c, v: pct(v) if ("Margen" in f or "%" in f) else
                      (num(v, 2) if "Cobertura" in f else (num(v, 0) if "días" in f else eur(v))))
            inf.figura(gr.fig_comparativo({ref: protocolos[ref], et: mc}, KPIS_GRAFICOS, 3),
                       f"Distribución Monte Carlo · {et} vs. {ref}", 17)
        else:
            inf.figura(gr.fig_histogramas(mc, KPIS_GRAFICOS[:8]), f"Distribución de los KPIs clave en las réplicas · {et}", 17)

    # --- Metodología
    avanza("Metodología")
    inf.salto(); inf.h("Anexo. Metodología e hipótesis de cálculo", 1)
    inf.vinetas([
        "Cada escenario se simula con el mismo motor de eventos discretos (SimPy); solo cambia su configuración.",
        "Tiempo en el sistema por paciente: T_sys = max(T_out) − min(T_in); tiempo de servicio T_ser = Σ T_server; "
        "tiempo en cola T_q = T_sys − T_ser (incluye la espera de Box, no registrada en T_wait).",
        f"Ventana de análisis: se descartan los primeros {ajustes['warmup']} días (arranque en vacío) y los últimos "
        f"{ajustes['cooldown']} (episodios incompletos).",
        "Utilización de camas = horas-cama ocupadas / (capacidad × horas de la ventana), incluido el bed-blocking. "
        "Longitud media de cola por la ley de Little.",
        "Costes de estancia: días de estancia clínica por la tarifa diaria. Costes fijos: los anuales se prorratean a la "
        "ventana; los únicos se imputan íntegramente.",
        "Réplicas con números aleatorios comunes: la réplica i de todos los escenarios usa la semilla inicial + i, lo que "
        "permite comparaciones pareadas (contraste t sobre las diferencias réplica a réplica, α = 5%).",
        "Un SLA se declara cumplido si lo cumple al menos el 95% de las réplicas, e incumplido si lo cumple como mucho el 5%.",
        "Los gráficos de tiempo en el sistema por perfil y de ocupación diaria de UCI corresponden a la primera réplica.",
    ])
    avanza("Guardando")
    return inf.guardar()
