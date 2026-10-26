# -*- coding: utf-8 -*-
"""Informe Word (.docx) del beneficio esperado por mixtura: teoría, simulación, estrategias y sensibilidad."""
import datetime
import io
import math

import numpy as np
import pandas as pd
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import mixtura as mx
from mixtura import num, pct, eur, fmt_kpi, fmt_mc, fmt_obj
import graficos as gr


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


ANCHO_UTIL_CM = 17.0          # A4 con márgenes de 2 cm


def _cm_por_caracter(pt):
    return 0.0205 * pt          # anchura media aproximada de un carácter de Calibri


def _anchos(cab, filas, pt):
    """Anchura de cada columna según su contenido; se reparte el ancho útil sin superar la página."""
    nat = []
    for j, c in enumerate(cab):
        palabras = max((len(w) for w in str(c).split()), default=1)
        # los textos largos (con espacios) pueden partirse en varias líneas; los números, no
        celdas = max((min(len(f[j]), 30) if " " in f[j].strip() and len(f[j]) > 30 else len(f[j]) for f in filas),
                     default=1)
        nat.append(max(palabras, celdas, 3) * _cm_por_caracter(pt) + 0.3)
    total = sum(nat)
    return nat, total


class Informe:
    """Envoltorio sencillo de python-docx con helpers para títulos, párrafos, tablas y figuras."""

    def __init__(self, titulo, subtitulo=""):
        self.doc = docx.Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Calibri"; estilo.font.size = Pt(10.5)
        self.n_tabla, self.n_fig = 0, 0
        self.doc.add_heading(titulo, level=0)
        if subtitulo:
            p = self.doc.add_paragraph(subtitulo); p.runs[0].italic = True

    def h(self, texto, nivel=1):
        self.doc.add_heading(texto, level=nivel)

    def p(self, texto, negrita_inicio=None):
        par = self.doc.add_paragraph()
        if negrita_inicio and texto.startswith(negrita_inicio):
            par.add_run(negrita_inicio).bold = True
            par.add_run(texto[len(negrita_inicio):])
        else:
            par.add_run(texto)
        par.paragraph_format.space_after = Pt(6)
        return par

    def vinetas(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def tabla(self, df, titulo, formato=None, indice=True, ancho_fuente=9, nombre_indice=""):
        """Tabla con anchos de columna calculados para que quepa siempre en el ancho de la página: si el contenido
        no cabe se reduce la letra (hasta 7 pt) y se reparten las columnas proporcionalmente."""
        self.n_tabla += 1
        cap = self.doc.add_paragraph()
        cap.add_run(f"Tabla {self.n_tabla}. {titulo}").bold = True
        cap.paragraph_format.keep_with_next = True
        cols = list(df.columns)
        cab = ([nombre_indice or (df.index.name or "")] if indice else []) + [str(c) for c in cols]
        filas = []
        for idx, fila in df.iterrows():
            filas.append(([num(idx, 0) if isinstance(idx, (int, np.integer)) else str(idx)] if indice else []) +
                         [formato(idx, c, fila[c]) if formato else _fmt(fila[c]) for c in cols])
        pt = ancho_fuente
        nat, total = _anchos(cab, filas, pt)
        while total > ANCHO_UTIL_CM and pt > 7:
            pt -= 0.5
            nat, total = _anchos(cab, filas, pt)
        anchos = [w * ANCHO_UTIL_CM / total for w in nat] if total > ANCHO_UTIL_CM else \
                 [w * min(ANCHO_UTIL_CM / total, 1.6) for w in nat]
        t = self.doc.add_table(rows=1, cols=len(cab))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        # la fila de cabecera se repite si la tabla continúa en otra página
        trPr = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trPr.append(rep)
        for j, col in enumerate(t.columns):          # rejilla de la tabla (la que usan Word y LibreOffice)
            col.width = Cm(anchos[j])
        for j, c in enumerate(cab):
            cel = t.rows[0].cells[j]; cel.text = ""; cel.width = Cm(anchos[j])
            r = cel.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(pt)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if j > 0 or not indice:
                cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _sombrear(cel, "2F5496")
        for i, vals in enumerate(filas):
            celdas = t.add_row().cells
            for j, v in enumerate(vals):
                celdas[j].text = ""; celdas[j].width = Cm(anchos[j])
                r = celdas[j].paragraphs[0].add_run(v); r.font.size = Pt(pt)
                if j > 0 or not indice:
                    celdas[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i % 2 == 1:
                    _sombrear(celdas[j], "EAF0F8")
        self.doc.add_paragraph()

    def figura(self, fig, titulo, ancho_cm=16.5):
        self.n_fig += 1
        self.doc.add_picture(gr.fig_a_png(fig), width=Cm(ancho_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figura {self.n_fig}. {titulo}"); r.italic = True; r.font.size = Pt(9)

    def salto(self):
        self.doc.add_page_break()

    def guardar(self):
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def _fmt(v):
    if isinstance(v, (bool, np.bool_)):
        return "Sí" if v else "No"
    if v is None:
        return "—"
    if isinstance(v, (int, np.integer)):
        return num(v, 0)
    if isinstance(v, (float, np.floating)):
        return "—" if math.isnan(v) else num(v, 3)
    return str(v)


def fmt_tabla(fila, col, v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "—"
    if "€" in col:
        return eur(v, 2)
    if col.startswith(("Teórica", "Simulada")) or col.startswith("P("):
        return pct(v, 2)
    return str(v)


def fmt_sens_factory(metrica):
    f = mx.fmt_metrica(metrica)
    def fmt(fila, col, v):
        if col == "Error relativo":
            return pct(v, 3)
        if metrica == "proporcion" and col.startswith("σ"):
            return pct(v, 3)
        if col.startswith("σ"):
            return eur(v, 3)
        return f(v)
    return fmt


def tabla_sens(df):
    return df.drop(columns=["Medias por réplica"])


def generar_informe(res, comunes, ajustes, titulo="Beneficio esperado por mixtura", sens=None):
    inf = Informe(titulo, f"Informe generado automáticamente el {datetime.datetime.now():%d/%m/%Y %H:%M} · "
                          "mixtura en dos pasos, simulación Monte Carlo y comparación de estrategias comerciales")
    V = comunes["visitantes"]
    # 1. Planteamiento
    inf.h("1. Planteamiento y parámetros", 1)
    inf.p("Cada cliente que entra en la tienda compra un horno (T = 1) con probabilidad Pr(compra) o se va sin "
          "comprar (T = 0, beneficio 0 €). Si compra, elige uno de los modelos con las probabilidades condicionadas "
          "pᵢ = P(modelo i | compra), y la tienda gana bᵢ. El beneficio por cliente B es, por tanto, una mixtura: "
          "B = 0 con probabilidad 1 − Pr(compra) y B = bᵢ con probabilidad Pr(compra)·pᵢ. Se simula en dos pasos: "
          "U₁ ~ Unif(0, 1) decide si compra (U₁ < Pr(compra)) y, solo si compra, U₂ ~ Unif(0, 1) elige el modelo "
          "por transformada inversa. E(B | compra) = Σ pᵢ·bᵢ y E(B) = Pr(compra)·E(B | compra). El beneficio neto "
          "descuenta el coste comercial por cliente que entra de cada estrategia.")
    esc = pd.DataFrame({n: {"P(compra)": pct(r["teoria"]["pc"], 0),
                            "Modelos (ganancia)": " / ".join(f"{m} {eur(b, 0)}" for m, b in zip(r["teoria"]["modelos"], r["teoria"]["b"])),
                            "P(modelo | compra)": " / ".join(pct(p, 0) for p in r["teoria"]["p"]),
                            "Coste por cliente": eur(r["teoria"]["coste"], 2)} for n, r in res.items()})
    inf.tabla(esc, "Estrategias comerciales (escenarios) analizadas", nombre_indice="Parámetro", formato=lambda f, c, v: v)
    desc = [f"{n}: {r['esc']['descripcion']}" for n, r in res.items() if r["esc"].get("descripcion")]
    if desc:
        inf.vinetas(desc)
    ob = comunes["objetivos"]
    def o(k, f):
        return f(ob[k]) if ob.get(k) is not None else "no se evalúa"
    com = pd.DataFrame({"Valor": {
        "Clientes que entran por día": num(V, 0),
        "Objetivo · beneficio neto mínimo por cliente": o("BNETO_MIN", lambda v: eur(v, 2)),
        "Objetivo · proporción mínima que compra": o("PCOMPRA_MIN", lambda v: pct(v, 1)),
        "Objetivo · beneficio mínimo por comprador": o("BCOMPRA_MIN", lambda v: eur(v, 2)),
        "Objetivo · compras de gama alta mínimas": o("ALTA_MIN", lambda v: pct(v, 1)),
        "Simulación · clientes por réplica (nsim)": num(ajustes["nsim"], 0),
        "Simulación · réplicas (semillas)": f"{ajustes['replicas']} (semillas {ajustes['semilla']} … "
                                            f"{int(ajustes['semilla']) + int(ajustes['replicas']) - 1})"}})
    inf.tabla(com, "Parámetros comunes, objetivos y ajustes de simulación", nombre_indice="Parámetro",
              formato=lambda f, c, v: v)

    # 2. Síntesis
    inf.h("2. Síntesis y recomendación", 1)
    for n, r in res.items():
        inf.p(mx.texto_escenario(n, r), negrita_inicio=n + ":")
    inf.p("Recomendación: " + mx.texto_recomendacion(res), negrita_inicio="Recomendación:")
    if len(res) > 1:
        inf.tabla(mx.tabla_comparativa(res), "Comparación de estrategias (filas «Teórico»: fórmulas de la mixtura; "
                  "filas «Simulado»: media de las réplicas)", formato=fmt_kpi, nombre_indice="KPI")
        inf.tabla(mx.tabla_objetivos(res), "Cumplimiento de los objetivos (valores teóricos)", nombre_indice="Objetivo",
                  formato=lambda f, c, v: str(v))
        inf.figura(gr.fig_comparativa(res), "Proporción que compra, beneficio por comprador, beneficio neto por cliente "
                                            "y beneficio neto diario por estrategia (teóricos)")
        inf.figura(gr.fig_validacion(res), "Error relativo de la simulación frente a los valores teóricos")

    # 3. Detalle por estrategia
    for k, (n, r) in enumerate(res.items(), start=3):
        inf.salto()
        t, mc = r["teoria"], r["mc"]
        inf.h(f"{k}. Estrategia «{n}»", 1)
        if r["esc"].get("descripcion"):
            inf.p(r["esc"]["descripcion"])
        inf.h(f"{k}.1 Resultados teóricos", 2)
        teo = pd.DataFrame({"Valor": {
            "Proporción que compra Pr(compra)": pct(t["pc"], 2),
            "Beneficio medio por comprador E(B | compra) = Σ pᵢ·bᵢ": eur(t["E_compra"]),
            "Beneficio medio por cliente E(B) = Pr(compra)·E(B | compra)": eur(t["E"]),
            "Pérdida de beneficio medio por los que no compran": eur(t["coste_no_compra"]),
            "Desviación típica por cliente σ(B)": eur(t["sd"]),
            "Coste comercial por cliente": eur(t["coste"]),
            "Beneficio neto por cliente": eur(t["E_neto"]),
            f"Beneficio neto diario ({num(V, 0)} clientes)": eur(t["dia_neto"], 0),
            "Compras de gama alta (todos salvo el 1.er modelo)": pct(t["gama_alta"], 2)}})
        inf.tabla(teo, f"KPIs teóricos · {n}", nombre_indice="Medida", formato=lambda f, c, v: v)
        inf.tabla(r["objetivos"], f"Objetivos (valores teóricos) · {n}", formato=fmt_obj, nombre_indice="Objetivo")
        inf.figura(gr.fig_descomposicion(t, f"Descomposición del beneficio esperado · {n}"),
                   f"Contribución de cada desenlace a E(B) · {n}")

        inf.h(f"{k}.2 Resultados simulados", 2)
        inf.p(f"{ajustes['replicas']} réplicas independientes de {num(ajustes['nsim'], 0)} clientes simulados en dos "
              "pasos. Cada KPI se calcula en cada réplica; la tabla muestra la media de las réplicas y su IC95 % (t de "
              "Student). Todas las estrategias usan las mismas semillas (números aleatorios comunes).")
        inf.tabla(r["resumen_mc"], f"KPIs simulados frente a teóricos · {n}", formato=fmt_mc, nombre_indice="KPI")
        inf.figura(gr.fig_replicas(mc, t, f"Beneficio medio en cada réplica · {n}"),
                   f"E(B) y E(B | compra) de cada réplica con su IC95 % y valores teóricos · {n}")

        inf.h(f"{k}.3 Distribución del beneficio por cliente (teórica y simulada)", 2)
        inf.tabla(mx.tabla_desenlaces(t, mc), f"Desenlaces de la mixtura · {n}", formato=fmt_tabla, nombre_indice="Desenlace")
        inf.tabla(mx.tabla_modelos(t, mc), f"Modelos entre los compradores e intervalo de U₂ · {n}", formato=fmt_tabla,
                  nombre_indice="Modelo")
        inf.figura(gr.fig_desenlaces(t, mc, f"Distribución del beneficio por cliente · {n}"),
                   f"Probabilidad de cada desenlace: teórica (barras) y simulada (rombos) · {n}")

        inf.h(f"{k}.4 Funcionamiento simulado", 2)
        tr = mc["traza"]
        u1 = float(tr["U1"].iloc[0]); u2 = float(tr["U2"].iloc[0]) if tr["Compra"].iloc[0] else 0.5
        inf.figura(gr.fig_mecanismo(t, u1, u2, f"Mixtura en dos pasos · primer cliente · {n}"),
                   f"Paso 1 (¿compra?) y paso 2 (qué modelo) para el primer cliente simulado (U₁ = {num(u1, 2)}) · {n}")
        inf.figura(gr.fig_traza(tr, t, f"Primeros {len(tr)} clientes · {n}"),
                   f"U₁ de los primeros {len(tr)} clientes y modelo elegido por los que compran · {n}")

    if sens is not None:
        inf.salto()
        df, cfg = sens["df"], sens["cfg"]
        inf.h(f"{len(res) + 3}. Sensibilidad del estimador Monte Carlo", 1)
        inf.p(f"Se varía nsim entre {num(df.index.min(), 0)} y {num(df.index.max(), 0)} ({len(df)} puntos en escala "
              f"logarítmica) con {cfg['replicas']} réplicas por punto (semillas {cfg['semilla']} … "
              f"{cfg['semilla'] + cfg['replicas'] - 1}). " + sens["texto"])
        inf.figura(gr.fig_sens_convergencia(df, cfg["metrica"]), "Convergencia de la estimación hacia el valor teórico")
        inf.figura(gr.fig_sens_variabilidad(df, cfg["metrica"]), "Variabilidad de la estimación entre réplicas y σ teórica del estimador")
        inf.tabla(tabla_sens(df), "Barrido del número de clientes simulados", formato=fmt_sens_factory(cfg["metrica"]),
                  nombre_indice="nsim")

    inf.salto()
    inf.h("Anexo. Metodología", 1)
    inf.vinetas([
        "Paso 1: U₁ ~ Unif(0, 1); T = 1 (compra) si U₁ < Pr(compra) y T = 0 en otro caso (B = 0 €).",
        "Paso 2 (solo si T = 1): U₂ ~ Unif(0, 1); modelo = menor i tal que U₂ ≤ F(i) = p₁ + … + pᵢ; B = bᵢ.",
        "E(B | compra) = Σ pᵢ·bᵢ ; E(B) = Pr(compra)·E(B | compra) ; Var(B) = Pr(compra)·Σ pᵢ·bᵢ² − E(B)².",
        "Beneficio neto por cliente = E(B) − coste comercial ; diario = clientes por día × beneficio neto.",
        "σ teórica de los estimadores con nsim clientes: σ(B)/√nsim (todos), √(Pr(compra)(1 − Pr(compra))/nsim) "
        "(proporción) y σ(B | compra)/√(nsim·Pr(compra)) (compradores, aproximada).",
        "IC95 %: media de las réplicas ± t₀.₉₇₅·s/√R, con s la desviación típica entre réplicas.",
        "Basado en el Ejemplo 4 de «B-70. Distribuciones derivadas», Javier Morales y Asun Martínez Mayoral, equipo "
        "IA4LEGOS (CC BY-SA 4.0).",
    ])
    return inf.guardar()
