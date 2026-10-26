# -*- coding: utf-8 -*-
"""Informe Word (.docx) del sistema M/M/s/K: teoría, validación por simulación, costes y comparación de escenarios."""
import datetime
import io
import math

import numpy as np
import pandas as pd
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import mmsk
from mmsk import num, pct, eur, fmt_kpi, unidad_larga, unidad_sing
import graficos as gr


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


class Informe:
    """Envoltorio sencillo de python-docx con helpers para títulos, párrafos, tablas y figuras."""

    def __init__(self, titulo, subtitulo=""):
        self.doc = docx.Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Calibri"; estilo.font.size = Pt(10.5)
        self.n_tabla, self.n_fig = 0, 0
        self.doc.add_heading(titulo, level=0)
        if subtitulo:
            p = self.doc.add_paragraph(subtitulo); p.runs[0].italic = True

    def h(self, texto, nivel=1):
        self.doc.add_heading(texto, level=nivel)

    def p(self, texto, negrita_inicio=None):
        par = self.doc.add_paragraph()
        if negrita_inicio == "auto":
            i = texto.find(": ")
            negrita_inicio = texto[:i + 1] if 0 < i < 70 else None
        if negrita_inicio and texto.startswith(negrita_inicio):
            par.add_run(negrita_inicio).bold = True
            par.add_run(texto[len(negrita_inicio):])
        else:
            par.add_run(texto)
        par.paragraph_format.space_after = Pt(6)
        return par

    def vinetas(self, items):
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def tabla(self, df, titulo, formato=None, indice=True, ancho_fuente=8.5, nombre_indice=""):
        self.n_tabla += 1
        cap = self.doc.add_paragraph()
        cap.add_run(f"Tabla {self.n_tabla}. {titulo}").bold = True
        cols = list(df.columns)
        t = self.doc.add_table(rows=1, cols=len(cols) + (1 if indice else 0))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cab = ([nombre_indice or (df.index.name or "")] if indice else []) + [str(c) for c in cols]
        for j, c in enumerate(cab):
            cel = t.rows[0].cells[j]; cel.text = ""
            r = cel.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(ancho_fuente)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _sombrear(cel, "2F5496")
        for i, (idx, fila) in enumerate(df.iterrows()):
            celdas = t.add_row().cells
            vals = ([str(idx)] if indice else []) + [formato(idx, c, fila[c]) if formato else _fmt(fila[c]) for c in cols]
            for j, v in enumerate(vals):
                celdas[j].text = ""
                r = celdas[j].paragraphs[0].add_run(v); r.font.size = Pt(ancho_fuente)
                if j > 0 or not indice:
                    celdas[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i % 2 == 1:
                    _sombrear(celdas[j], "EAF0F8")
        self.doc.add_paragraph()

    def figura(self, fig, titulo, ancho_cm=16.5):
        self.n_fig += 1
        self.doc.add_picture(gr.fig_a_png(fig), width=Cm(ancho_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figura {self.n_fig}. {titulo}"); r.italic = True; r.font.size = Pt(9)

    def salto(self):
        self.doc.add_page_break()

    def guardar(self):
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def _fmt(v):
    if isinstance(v, (bool, np.bool_)):
        return "Sí" if v else "No"
    if v is None:
        return "—"
    if isinstance(v, (int, np.integer)):
        return num(v, 0)
    if isinstance(v, (float, np.floating)):
        if math.isnan(v):
            return "—"
        if math.isinf(v):
            return "∞"
        return num(v, 3)
    return str(v)


def fmt_mc(fila, col, v):
    if col == "Teórico dentro del IC95":
        return "—" if v is None else ("Sí" if v else "No")
    if col == "Error relativo":
        return "—" if v is None or (isinstance(v, float) and math.isnan(v)) else pct(v, 2)
    if fila in ("P(espera)", "P(bloqueo)"):
        return pct(v, 2)
    return num(v, 4)


def fmt_obj(fila, col, v):
    if col == "Valor":
        if "ρ" in fila or "Probabilidad" in fila:
            return pct(v, 2)
        return num(v, 4)
    return str(v)


def fmt_barrido(fila, col, v):
    if col in ("ρ", "ρ̄"):
        return pct(v, 1)
    if col == "P(bloqueo)":
        return pct(v, 2)
    if col in ("CTS", "CTSn", "CTQ", "CT"):
        return eur(v)
    return num(v, 3)


def tabla_teoria(t, cst, unidad):
    u = unidad
    us = unidad_sing(u)
    filas = [("a = λ/μ (carga ofrecida, Erlangs)", num(t["a"], 4)), ("ρ = λ/(sμ) (carga nominal por servidor)", pct(t["rho"], 2)),
             ("Régimen", t["estado"]), ("K (capacidad) · cola máxima K − s", f"{t['K']} · {t['K'] - t['s']}"),
             ("P0 (sistema vacío)", pct(t["P0"], 2)), ("P(bloqueo) = p_K (llegadas rechazadas)", pct(t["Pbloqueo"], 2)),
             (f"λ̄ = λ(1 − p_K) (clientes aceptados/{us})", num(t["lambda_bar"], 4)),
             ("ρ̄ = λ̄/(sμ) (ocupación efectiva por servidor)", pct(t["rho_eff"], 2)),
             ("P(espera) de un cliente aceptado", pct(t["Pespera"], 2)), ("L (clientes en el sistema)", num(t["L"], 4)),
             ("Lq (clientes en cola)", num(t["Lq"], 4)), (f"W (tiempo en el sistema, {unidad_larga(u)})", num(t["W"], 4)),
             (f"Wq (tiempo en cola, {unidad_larga(u)})", num(t["Wq"], 4)), ("Servidores ocupados (L − Lq)", num(t["Ocupados"], 4)),
             (f"c_t = c_q + c_s (€/{us})", eur(cst["c_t"])), (f"CTS (€/{us})", eur(cst["CTS"])),
             (f"CTSn (€/{us})", eur(cst["CTSn"])), (f"CTQ (€/{us})", eur(cst["CTQ"])),
             (f"CT (€/{us})", eur(cst["CT"]))]
    return pd.DataFrame(filas, columns=["Medida", "Valor"]).set_index("Medida")


def tabla_probs_informe(t, pn_mc, max_filas=41):
    """Tabla de probabilidad de n clientes, teórica y simulada, formateada para el informe."""
    df = mmsk.tabla_probabilidades(t, pn_mc).head(max_filas)
    df = df.drop(columns=["Clientes en cola", "Servidores ocupados"])
    df["Región"] = df["Región"].str.split(" ").str[0].str.replace("lleno", "lleno (K)")
    for c in df.columns:
        if c.startswith(("Teórica", "Simulada")):
            df[c] = df[c].map(lambda v: pct(v, 2))
    return df


def generar_informe(res, comunes, ajustes, titulo="Sistema de colas M/M/s/K", ref=None):
    u = comunes["unidad"]
    us, ul = unidad_sing(u), unidad_larga(u)
    inf = Informe(titulo, f"Informe generado automáticamente el {datetime.datetime.now():%d/%m/%Y %H:%M} · "
                          f"teoría M/M/s/K, simulación de eventos discretos y análisis de costes")
    # 1. Planteamiento
    inf.h("1. Planteamiento y parámetros", 1)
    inf.p("Sistema de colas M/M/s/K: llegadas de Poisson a tasa λ, tiempos de servicio exponenciales de tasa μ por "
          "servidor, s servidores en paralelo con una única cola FIFO y capacidad máxima de K clientes en el sistema (en "
          "servicio y en cola). Una llegada que encuentra el sistema lleno (N = K) se pierde. El proceso N(t) es de "
          "nacimiento-muerte con λ_n = λ para n < K, λ_K = 0 y μ_n = min(n, s)·μ; al ser finito, SIEMPRE tiene régimen "
          "estacionario, incluso con ρ = λ/(sμ) ≥ 100 % (régimen saturado). La tasa efectiva de entrada es "
          f"λ̄ = λ(1 − p_K). Todas las tasas y tiempos se expresan en {unidad_larga(u)}.")
    esc = pd.DataFrame({n: {"λ (llegadas)": r["esc"]["lambda"], "μ (servicio por servidor)": r["esc"]["mu"],
                            "s (servidores)": r["esc"]["s"], "K (capacidad)": r["esc"]["K"], f"c_s (€/{us}·servidor)": r["esc"]["cs"],
                            "Descripción": r["esc"].get("descripcion", "") or "—"} for n, r in res.items()})
    inf.tabla(esc, "Escenarios analizados", nombre_indice="Parámetro",
              formato=lambda f, c, v: v if isinstance(v, str) else num(v, 3 if "λ" in f or "μ" in f else (0 if f.startswith(("s ", "K ")) else 2)))
    ob = comunes["objetivos"]
    com = pd.DataFrame({"Valor": {
        f"Coste de espera c_q (€/ut por cliente; ut = {us})": eur(comunes["cq"]),
        f"Coste por servidor c_s por defecto (€/ut; ut = {us})": eur(comunes["cs"]),
        "s máximo explorado en el dimensionamiento (≤ K)": num(comunes["s_max"], 0),
        "Objetivo · ρ̄ máxima (ocupación efectiva)": pct(ob["RHO_MAX"], 0) if ob.get("RHO_MAX") is not None else "no se evalúa",
        f"Objetivo · Wq máxima ({ul})": num(ob["WQ_MAX"], 3) if ob.get("WQ_MAX") is not None else "no se evalúa",
        "Objetivo · P(bloqueo) máxima": pct(ob["PBLOQ_MAX"], 1) if ob.get("PBLOQ_MAX") is not None else "no se evalúa",
        "Objetivo · Lq máxima": num(ob["LQ_MAX"], 2) if ob.get("LQ_MAX") is not None else "no se evalúa",
        f"Simulación · duración de cada réplica ({ul})": num(ajustes["duracion"], 0),
        "Simulación · réplicas (semillas)": f"{ajustes['replicas']} (semillas {ajustes['semilla']} … "
                                            f"{int(ajustes['semilla']) + int(ajustes['replicas']) - 1})",
        "Diagrama de funcionamiento": f"primeras {mmsk.N_DIAGRAMA} llegadas"}})
    inf.tabla(com, "Parámetros comunes, objetivos de servicio y ajustes de simulación", nombre_indice="Parámetro",
              formato=lambda f, c, v: v)

    # 2. Síntesis
    inf.h("2. Síntesis y recomendación", 1)
    for n, r in res.items():
        inf.p(mmsk.texto_escenario(n, r, u), negrita_inicio=n + ":")
    inf.p("Recomendación: " + mmsk.texto_recomendacion(res, u), negrita_inicio="Recomendación:")
    if len(res) > 1:
        comp = mmsk.tabla_comparativa(res)
        inf.tabla(comp, "Comparación de escenarios (filas «Teórico»: fórmulas del modelo; filas «Simulado»: media de las "
                        "réplicas Monte Carlo)", formato=fmt_kpi, nombre_indice="KPI",
                  ancho_fuente=7.5)
        inf.tabla(mmsk.tabla_objetivos(res), "Cumplimiento de los objetivos de servicio (valores teóricos)", nombre_indice="Objetivo",
                  formato=lambda f, c, v: str(v), ancho_fuente=8)
        inf.figura(gr.fig_comparativa(res, ul), "Ocupación efectiva, probabilidad de bloqueo, tiempo medio en cola y coste total por escenario (teóricos)")
        inf.figura(gr.fig_validacion(res), "Error relativo de la simulación frente a los valores teóricos")

    # 3. Detalle por escenario
    for k, (n, r) in enumerate(res.items(), start=3):
        inf.salto()
        t, cst = r["teoria"], r["costes"]
        inf.h(f"{k}. Escenario «{n}» · M/M/{t['s']}/{t['K']}", 1)
        if r["esc"].get("descripcion"):
            inf.p(r["esc"]["descripcion"])
        inf.h(f"{k}.1 Resultados teóricos: distribución límite y KPIs", 2)
        inf.p("Valores exactos calculados con las fórmulas del modelo M/M/s/K (distribución límite truncada en K).")
        inf.tabla(tabla_teoria(t, cst, u), f"KPIs teóricos y costes · {n}", nombre_indice="Medida", formato=lambda f, c, v: v)
        inf.figura(gr.fig_pn(t, f"Probabilidad teórica de n clientes en el sistema · {n}"),
                   f"Probabilidad teórica de n clientes en el sistema, p_n = P(N = n) · {n}")
        inf.tabla(r["objetivos"], f"Objetivos de servicio (valores teóricos) · {n}", formato=fmt_obj, nombre_indice="Objetivo")

        inf.h(f"{k}.2 Resultados simulados: validación por simulación de eventos discretos", 2)
        inf.p(f"{ajustes['replicas']} réplicas independientes de {num(ajustes['duracion'], 0)} {ul} cada una (unas "
              f"{num(r['mc']['replicas']['Llegadas'].mean(), 0)} llegadas medidas por réplica, de las que se rechazan unas {num(r['mc']['replicas']['Rechazos'].mean(), 0)}, tras descartar el primer {mmsk.CALENTAMIENTO:.0%} como calentamiento; relojes "
              "exponenciales exactos, sin discretizar el tiempo). L y Lq se estiman como promedios en el tiempo; P(bloqueo) "
              "como la fracción de llegadas rechazadas; λ̄ como aceptados por unidad de tiempo; W = L/λ̄ y Wq = Lq/λ̄ "
              "por la ley de Little. El IC95 % refleja la variabilidad entre réplicas.")
        inf.tabla(r["resumen_mc"], f"KPIs simulados frente a teóricos · {n}", formato=fmt_mc, nombre_indice="KPI", ancho_fuente=8)
        inf.h(f"{k}.3 Probabilidad de n clientes en el sistema (teórica y simulada)", 2)
        inf.p("P(N = n): probabilidad de que haya exactamente n clientes en el sistema; P(N ≤ n), su acumulada; P(N > n), la "
              "probabilidad de superar n clientes; p_K es la probabilidad de bloqueo. La teórica procede de las fórmulas; la simulada es la fracción del tiempo "
              "simulado (todas las réplicas) en que había n clientes.")
        inf.figura(gr.fig_pn_comparada(r["mc"]["pn"], t, f"Probabilidad de n clientes: simulada y teórica · {n}"),
                   f"Probabilidad de n clientes en el sistema: simulada (barras) frente a teórica (línea) · {n}")
        inf.tabla(tabla_probs_informe(t, r["mc"]["pn"]), f"Probabilidad de n clientes en el sistema · {n}", nombre_indice="n",
                  formato=lambda f, c, v: v, ancho_fuente=7.5)
        inf.h(f"{k}.4 Funcionamiento simulado", 2)
        inf.figura(gr.fig_traza(r["mc"]["traza"], t["s"], ul, f"Trayectoria N(t) simulada · {n}", K=t["K"]),
                   f"Trayectoria N(t) de la primera réplica (primeros 400 eventos) · {n}")
        inf.figura(gr.fig_diagrama(r["clientes"], t["s"], None, ul,
                                   f"Diagrama de funcionamiento simulado · {n}"),
                   f"Llegada, espera, servicio y salida de las primeras {len(r['clientes'])} llegadas "
                   f"({int(r['clientes']['Rechazado'].sum())} rechazadas, marcadas con ✕) · {n}")

        inf.h(f"{k}.5 Costes y dimensionamiento (teóricos)", 2)
        inf.p("CTS: coste de los servidores ocupados (a c_t = c_q + c_s). CTSn: coste de los servidores desocupados (a c_s). "
              "CTQ: coste de espera de los clientes en cola (λ·c_q·Wq, con λ nominal). CT = CTS + CTSn + CTQ. El coste de los "
              f"clientes rechazados no se incluye. El dimensionamiento explora s = 1 … min(s máx, K) con K = {t['K']} fijo.")
        inf.figura(gr.fig_costes(cst, us, f"Desglose de costes · {n}"), f"Desglose del coste total · {n}")
        inf.figura(gr.fig_barrido(r["barrido"], t["s"], r["s_opt"], us, f"Coste total frente a s · {n}"),
                   f"Coste total y probabilidad de bloqueo frente al número de servidores · {n}")
        inf.tabla(r["barrido"], f"Barrido del número de servidores · {n}", formato=fmt_barrido, nombre_indice="s",
                  ancho_fuente=7.5)

    inf.salto()
    inf.h("Anexo. Metodología", 1)
    inf.vinetas([
        "P0 = [Σ_{n<s} aⁿ/n! + aˢ/s!·(1 − ρ^{K−s+1})/(1 − ρ)]⁻¹ (si ρ = 1, el último factor vale K − s + 1).",
        "p_n = P0·aⁿ/n! si n < s ; p_n = P0·aⁿ/(s!·s^{n−s}) si s ≤ n ≤ K ; P(bloqueo) = p_K.",
        "L = Σ n·p_n ; λ̄ = λ(1 − p_K) ; Lq = L − λ̄/μ ; W = L/λ̄ ; Wq = W − 1/μ (ley de Little con la tasa efectiva).",
        "ρ̄ = λ̄/(sμ) (ocupación efectiva) ; P(espera) de un aceptado = Σ_{s≤n<K} p_n/(1 − p_K).",
        "Simulación: proceso de nacimiento-muerte con relojes exponenciales competitivos; una llegada con N = K se rechaza; p_n simulada como fracción "
        "del tiempo total en cada estado, combinando las réplicas ponderadas por su duración.",
        "Costes por unidad de tiempo: CTS = (c_q + c_s)·(L − Lq) ; CTSn = c_s·(s − (L − Lq)) ; CTQ = λ·c_q·Wq.",
        "Basado en «D-190. Sistemas de Colas», Javier Morales y Asun Martínez Mayoral, equipo IA4LEGOS (CC BY-SA 4.0).",
    ])
    return inf.guardar()
